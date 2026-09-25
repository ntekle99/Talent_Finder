#!/usr/bin/env python3
"""Stream every LCA xlsx in data/raw/, aggregate certified hires for our ~50 tech
employers into data/panel/talent_panel.csv (company_id x month).

Low-memory: openpyxl read_only streams rows; we keep only aggregated cells for matched
companies. Parallel across files. Prints a per-company coverage report so we can sanity
-check the employer matching before trusting the backtest.
"""
from __future__ import annotations
import csv, sys, os, time
from collections import defaultdict
from datetime import datetime, date
from multiprocessing import Pool
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))
import openpyxl
from talent_finder.ingest.h1b import _ALIASES, _pick, seniority_of, _to_float, _to_month
from company_map import Matcher, COMPANIES

RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "panel"
SKIP = {"H-1B_Case_Data_FY2008.xlsx", "H-1B_Case_Data_FY2009.xlsx", "Icert_LCA_FY2009.xlsx"}


def _month_of(v) -> str | None:
    if isinstance(v, (datetime, date)):
        return f"{v.year:04d}-{v.month:02d}"
    return _to_month(str(v)) if v is not None else None


def aggregate_file(path_str: str) -> dict:
    """Worker: -> {'cells': {(cid,month): [hires,senior_sum,wage_sum,wage_n]},
                    'file': name, 'scanned': int, 'matched': int}."""
    path = Path(path_str)
    matcher = Matcher(COMPANIES)
    cells: dict[tuple[str, str], list] = defaultdict(lambda: [0, 0.0, 0.0, 0])
    scanned = matched = 0
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(rows)]
        ci = {k: _pick(header, v) for k, v in _ALIASES.items()}
        if ci["employer"] is None or ci["date"] is None:
            wb.close()
            return {"cells": {}, "file": path.name, "scanned": 0, "matched": 0,
                    "note": "no employer/date column"}
        for row in rows:
            scanned += 1
            st = ci["status"]
            if st is not None and st < len(row):
                s = row[st]
                if s is None or "CERTIF" not in str(s).upper():
                    continue
            emp = row[ci["employer"]] if ci["employer"] < len(row) else None
            cid = matcher.match(str(emp) if emp is not None else "")
            if cid is None:
                continue
            month = _month_of(row[ci["date"]]) if ci["date"] < len(row) else None
            if month is None:
                continue
            title = row[ci["title"]] if ci["title"] is not None and ci["title"] < len(row) else ""
            cell = cells[(cid, month)]
            cell[0] += 1
            cell[1] += seniority_of(str(title) if title else "")
            if ci["wage"] is not None and ci["wage"] < len(row):
                w = _to_float(str(row[ci["wage"]])) if row[ci["wage"]] is not None else None
                if w and 20_000 < w < 1_000_000:  # annual salaries only; drop hourly/garbage
                    cell[2] += w
                    cell[3] += 1
            matched += 1
        wb.close()
    except Exception as e:
        return {"cells": {}, "file": path.name, "scanned": scanned, "matched": matched,
                "note": f"ERROR {e}"}
    return {"cells": {k: v for k, v in cells.items()}, "file": path.name,
            "scanned": scanned, "matched": matched}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    files = sorted(str(p) for p in RAW.glob("*.xlsx") if p.name not in SKIP)
    print(f"parsing {len(files)} files (skipping {len(SKIP)} pre-2010 no-employer files)\n", flush=True)

    merged: dict[tuple[str, str], list] = defaultdict(lambda: [0, 0.0, 0.0, 0])
    per_company: dict[str, int] = defaultdict(int)
    t0 = time.time()
    nproc = min(8, os.cpu_count() or 4)
    with Pool(nproc) as pool:
        for res in pool.imap_unordered(aggregate_file, files):
            note = res.get("note", "")
            print(f"  {res['file']:42s} scanned={res['scanned']:>8} matched={res['matched']:>7} {note}", flush=True)
            for (cid, month), v in res["cells"].items():
                m = merged[(cid, month)]
                m[0] += v[0]; m[1] += v[1]; m[2] += v[2]; m[3] += v[3]
                per_company[cid] += v[0]

    # write panel (departures=0 until WARN ingest; source_mask=2 => H1B)
    path = OUT / "talent_panel.csv"
    from company_map import ticker_of
    tk = ticker_of()
    with path.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["company_id","ticker","month","hires","senior_hires","mean_wage","departures","source_mask"])
        for (cid, month) in sorted(merged.keys()):
            hires, senior_sum, wage_sum, wage_n = merged[(cid, month)]
            w.writerow([cid, tk.get(cid, ""), month, hires, round(senior_sum, 2),
                        int(wage_sum / wage_n) if wage_n else 0, 0, 2])

    print(f"\nwrote {len(merged)} panel rows -> {path}  ({time.time()-t0:.0f}s)")
    print("\ncoverage (total certified hires matched, FY2010-2026):")
    for co in COMPANIES:
        print(f"  {co.company_id:20s} {co.ticker:6s} {per_company.get(co.company_id, 0):>8}")
    zero = [co.company_id for co in COMPANIES if per_company.get(co.company_id, 0) == 0]
    if zero:
        print("\n  WARNING zero matches (check fragments):", ", ".join(zero))


if __name__ == "__main__":
    main()
