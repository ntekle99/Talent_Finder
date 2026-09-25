#!/usr/bin/env python3
"""Broad wage-premium talent signal over ALL public H1B sponsors (SEC-matched) ->
data/panel/talent_wage_broad.csv (ticker, year, hires, wage_ratio).

Same role-mix-adjusted wage_ratio as build_wage_talent.py, but the universe is every public
company in the H1B data (~hundreds), matched to tickers via SEC's official list, instead of 50
hand-picked names. This is the statistical-power upgrade for the backtest.
"""
from __future__ import annotations
import sys, csv, os, time
from collections import defaultdict
from datetime import datetime, date
from multiprocessing import Pool
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))
import openpyxl
from talent_finder.ingest.h1b import _pick, _to_float
from build_wage_talent import ALIASES, annualize, _year, SKIP
from sec_universe import SecMatcher

RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "panel" / "talent_wage_broad.csv"


def process(path_str: str) -> dict:
    path = Path(path_str)
    matcher = SecMatcher()
    nat = defaultdict(lambda: [0, 0.0])
    co = defaultdict(lambda: [0, 0.0])
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(rows)]
        ci = {k: _pick(header, v) for k, v in ALIASES.items()}
        if None in (ci["employer"], ci["soc"], ci["wage"], ci["date"]):
            wb.close(); return {"nat": {}, "co": {}, "file": path.name, "note": "missing cols"}
        for row in rows:
            st = ci["status"]
            if st is not None and st < len(row):
                s = row[st]
                if s is None or "CERTIF" not in str(s).upper():
                    continue
            if ci["wage"] >= len(row) or ci["soc"] >= len(row) or ci["date"] >= len(row):
                continue
            wage = _to_float(str(row[ci["wage"]])) if row[ci["wage"]] is not None else None
            if not wage:
                continue
            unit = row[ci["unit"]] if ci["unit"] is not None and ci["unit"] < len(row) else "Year"
            wage = annualize(wage, str(unit))
            if not (20_000 < wage < 2_000_000):
                continue
            year = _year(row[ci["date"]])
            soc = str(row[ci["soc"]])[:7].strip() if row[ci["soc"]] else ""
            if year is None or not soc:
                continue
            k = (soc, year)
            nat[k][0] += 1; nat[k][1] += wage
            tkr = matcher.match(str(row[ci["employer"]]) if row[ci["employer"]] is not None else "")
            if tkr:
                ck = (tkr, soc, year)
                co[ck][0] += 1; co[ck][1] += wage
        wb.close()
    except Exception as e:
        return {"nat": {}, "co": {}, "file": path.name, "note": f"ERR {e}"}
    return {"nat": dict(nat), "co": dict(co), "file": path.name}


def main():
    files = sorted(str(p) for p in RAW.glob("*.xlsx") if p.name not in SKIP)
    print(f"streaming {len(files)} files (broad universe)...\n", flush=True)
    nat = defaultdict(lambda: [0, 0.0]); co = defaultdict(lambda: [0, 0.0])
    t0 = time.time()
    with Pool(min(8, os.cpu_count() or 4)) as pool:
        for res in pool.imap_unordered(process, files):
            for k, v in res["nat"].items():
                nat[k][0] += v[0]; nat[k][1] += v[1]
            for k, v in res["co"].items():
                co[k][0] += v[0]; co[k][1] += v[1]
            print(f"  {res['file']:42s} {res.get('note','ok')}", flush=True)

    nat_mean = {k: v[1] / v[0] for k, v in nat.items() if v[0] > 0}
    agg = defaultdict(lambda: [0, 0.0, 0.0])
    for (tkr, soc, year), (n, s) in co.items():
        ref = nat_mean.get((soc, year))
        if ref:
            cell = agg[(tkr, year)]
            cell[0] += n; cell[1] += s; cell[2] += n * ref
    rows = []
    for (tkr, year), (hires, actual, expected) in sorted(agg.items()):
        if hires >= 20 and expected > 0:
            rows.append((tkr, year, hires, round(actual / expected, 4)))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="") as f:
        w = csv.writer(f); w.writerow(["ticker", "year", "hires", "wage_ratio"]); w.writerows(rows)
    tickers = sorted({t for t, _, _, _ in rows})
    print(f"\nwrote {len(rows)} ticker-year rows, {len(tickers)} distinct tickers "
          f"-> {OUT}  ({time.time()-t0:.0f}s)")
    (ROOT / "data" / "panel" / "broad_tickers.txt").write_text("\n".join(tickers))
    print("saved ticker universe -> data/panel/broad_tickers.txt")


if __name__ == "__main__":
    main()
